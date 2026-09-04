"""Office 文档（PPTX/DOCX）结构化解析：段落 + 表格 + 代码。"""

import re
from pathlib import Path

from app.services.parsing.blocks import BlockType, ParsedDocument, StructuredBlock
from app.services.parsing.code import detect_code_language, looks_like_code


def parse_office(path: Path) -> ParsedDocument:
    """解析 pptx/docx 为结构化块；旧版 .ppt/.doc 明确报错提示转格式。"""
    ext = path.suffix.lower()
    if ext not in {".pptx", ".docx"}:
        raise ValueError(
            f"旧版 {ext} 二进制格式暂不支持解析，请先用 Office 另存为"
            f"{'.pptx' if ext == '.ppt' else '.docx'} 后再上传。"
        )
    blocks: list[StructuredBlock] = []
    if ext == ".pptx":
        blocks = _parse_pptx(path)
    else:
        blocks = _parse_docx(path)
    return ParsedDocument(
        blocks=blocks,
        title=path.stem,
        source_path=str(path),
        parser="office",
    )


def _parse_pptx(path: Path) -> list[StructuredBlock]:
    from pptx import Presentation

    blocks: list[StructuredBlock] = []
    prs = Presentation(str(path))
    for slide_index, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.has_table:
                rows = []
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        rows.append("| " + " | ".join(cells) + " |")
                if rows:
                    blocks.append(
                        StructuredBlock(
                            block_type=BlockType.TABLE,
                            text="\n".join(rows),
                            page=slide_index,
                            source="office",
                        )
                    )
            elif shape.has_text_frame:
                text = "\n".join(
                    para.text.strip()
                    for para in shape.text_frame.paragraphs
                    if para.text.strip()
                )
                if not text:
                    continue
                blocks.append(_classify_text(text, slide_index, "office"))
    return blocks


def _parse_docx(path: Path) -> list[StructuredBlock]:
    from docx import Document

    blocks: list[StructuredBlock] = []
    doc = Document(str(path))
    for para in doc.paragraphs:
        if para.text.strip():
            blocks.append(_classify_text(para.text.strip(), None, "office"))
    for table in doc.tables:
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                rows.append("| " + " | ".join(cells) + " |")
        if rows:
            blocks.append(
                StructuredBlock(block_type=BlockType.TABLE, text="\n".join(rows), source="office")
            )
    return blocks


def _classify_text(text: str, page: int | None, source: str) -> StructuredBlock:
    """按内容把单段文本归类为标题/代码/公式/段落。"""
    from app.services.parsing.formula import FormulaExtractor

    if looks_like_code(text):
        return StructuredBlock(
            block_type=BlockType.CODE,
            text=text,
            page=page,
            language=detect_code_language(text) or "text",
            source=source,
        )
    if FormulaExtractor.looks_like_formula(text):
        return StructuredBlock(
            block_type=BlockType.FORMULA,
            text=text.strip("$ "),
            page=page,
            source=source,
        )
    if len(text) <= 60 and (
        (text.startswith("第") and ("章" in text[:8] or "节" in text[:8]))
        or re.match(r"^\d+(\.\d+)*[、.．\s]", text)
    ):
        return StructuredBlock(
            block_type=BlockType.HEADING, text=text, page=page, source=source
        )
    return StructuredBlock(
        block_type=BlockType.PARAGRAPH, text=text, page=page, source=source
    )
