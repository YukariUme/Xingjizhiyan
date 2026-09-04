"""多模态能力单元测试：视觉 Mock、Office 文档解析、答疑图片输入。"""

import base64
import pathlib
import tempfile

from app.database import SessionLocal
from app.models import User
from app.services.agent.learning import LearningAgent
from app.services.knowledge_service import extract_text_file
from app.services.llm.base import LLMService
from app.services.prompt import PromptService
from app.services.rag.factory import get_rag_service
from app.services.vision import MockVisionService, get_vision_service


def test_vision_mock_provider():
    """视觉工厂默认返回 Mock，且能给出稳定描述。"""
    service = get_vision_service()
    assert isinstance(service, MockVisionService)
    text = service.analyze_image(b"\x89PNG\r\n\x1a\n" + b"\x00" * 64, prompt="学生上传的代码截图")
    assert "图片内容" in text
    assert "代码" in text


def test_office_docx_extraction():
    """Word(.docx) 解析：段落 + 表格都能提取为可检索文本。"""
    from docx import Document

    with tempfile.TemporaryDirectory() as tmp:
        path = pathlib.Path(tmp) / "讲义.docx"
        doc = Document()
        doc.add_paragraph("操作系统第 4 章：进程同步")
        doc.add_paragraph("信号量支持 P 与 V 两种原子操作。")
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "生产者"
        table.rows[0].cells[1].text = "消费者"
        doc.save(str(path))
        text = extract_text_file(path)
    assert "进程同步" in text
    assert "信号量" in text
    assert "生产者 | 消费者" in text


def test_tutor_chat_accepts_image():
    """答疑支持图片输入：VisionService 转文字后进入 RAG + LLM 流程。"""

    class _EchoLLM(LLMService):
        name = "echo"

        def generate(self, prompt, system=None, temperature=0.7, max_tokens=None) -> str:
            return '{"answer": "已结合图片分析", "knowledge_points": ["信号量"]}'

        def chat(self, messages, temperature=0.7, max_tokens=None) -> str:
            return self.generate("")

        async def stream(self, messages):
            yield self.generate("")

    # 1x1 透明 PNG（最小合法图片字节）
    png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
    )
    db = SessionLocal()
    try:
        student = db.query(User).filter(User.role == "student").first()
        agent = LearningAgent(llm=_EchoLLM(), rag=get_rag_service(), prompts=PromptService())
        result = agent.run(
            db,
            student,
            {
                "task": "tutor",
                "message": "请看这张代码截图",
                "mode": "hint",
                "history": [],
                "image_base64": base64.b64encode(png).decode("ascii"),
            },
        )
        assert result["answer"] == "已结合图片分析"
    finally:
        db.close()
