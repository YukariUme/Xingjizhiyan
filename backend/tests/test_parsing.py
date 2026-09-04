"""复杂教材结构化解析管线测试：OCR 抽象、PDF 文字层/扫描层、Office、结构化切分。"""

import pathlib
import tempfile

import pytest

from app.services.ocr.base import OCRProvider
from app.services.ocr.factory import get_ocr_provider
from app.services.ocr.mock import MockOCRProvider
from app.services.parsing.blocks import BlockType, ParsedDocument, StructuredBlock
from app.services.parsing.pipeline import parse_document_file
from app.services.parsing.table import _html_to_markdown
from app.services.rag.sqlite import SQLiteRAGService


def test_ocr_factory_returns_provider():
    """OCR 工厂返回统一接口实现（rapid 或 mock 均可识别）。"""
    provider = get_ocr_provider()
    assert isinstance(provider, OCRProvider)
    import base64

    png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
    )
    result = provider.recognize(png)
    assert isinstance(result.text, str)
    assert provider.available() is True or provider.name == "mock"


def test_parse_text_layer_pdf():
    """有文字层 PDF：正文 + 标题被结构化，表格由 pdfplumber 提取。"""
    import fitz

    with tempfile.TemporaryDirectory() as tmp:
        path = pathlib.Path(tmp) / "textbook.pdf"
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "Chapter 2 Linear List", fontsize=16)
        page.insert_text((72, 110), "A linear list is a finite sequence of data elements.", fontsize=11)
        page.insert_text((72, 130), "Sequential list supports random access.", fontsize=11)
        doc.save(str(path))
        doc.close()

        parsed = parse_document_file(path)
        text = parsed.to_text()
        assert "Linear List" in text
        assert "random access" in text
        types = {b.block_type.value for b in parsed.blocks}
        assert "paragraph" in types
        assert "heading" in types
        assert parsed.parser in {"native-pdf", "native-text"}


def test_parse_scanned_pdf_with_mock_ocr(monkeypatch):
    """无文字层 PDF：布局不可用时整页 OCR（Mock），管线仍产出结构化块。"""
    import fitz

    # 用最小 pixmap 生成图片型 PDF
    with tempfile.TemporaryDirectory() as tmp:
        path = pathlib.Path(tmp) / "scanned.pdf"
        doc = fitz.open()
        page = doc.new_page()
        pix = fitz.Pixmap(fitz.csGRAY, fitz.IRect(0, 0, 300, 300), 0)
        pix.clear_with(220)
        page.insert_image(fitz.Rect(0, 0, 300, 300), pixmap=pix)
        doc.save(str(path))
        doc.close()

        # 布局引擎替换为“整页 text 区域”，避免模型下载
        from app.services.parsing import layout as layout_mod

        class _FakeLayout:
            def detect(self, image, page=None):
                return [layout_mod.LayoutRegion(label="text", box=[0, 0, 0, 0], score=1.0, page=page)]

        monkeypatch.setattr(layout_mod, "LayoutService", _FakeLayout)
        parsed = parse_document_file(path, ocr_provider=MockOCRProvider())
        assert parsed.blocks
        assert any(b.source == "mock" for b in parsed.blocks)


def test_parse_docx_structured():
    """DOCX：段落 + 表格 → 结构化块（TABLE 类型）。"""
    from docx import Document

    with tempfile.TemporaryDirectory() as tmp:
        path = pathlib.Path(tmp) / "讲义.docx"
        doc = Document()
        doc.add_paragraph("操作系统第 4 章：进程同步")
        doc.add_paragraph("信号量支持 P 与 V 两种原子操作。")
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "生产者"
        table.rows[0].cells[1].text = "消费者"
        doc.save(str(path))
        parsed = parse_document_file(path)
        table_blocks = [b for b in parsed.blocks if b.block_type == BlockType.TABLE]
        assert table_blocks and "生产者 | 消费者" in table_blocks[0].text
        assert parsed.parser == "office"


def test_structured_chunking_keeps_blocks():
    """结构化块入库：表格/代码/公式保整，元数据携带 block_type 与页码。"""
    service = SQLiteRAGService()
    blocks = [
        StructuredBlock(block_type=BlockType.HEADING, text="第 4 章 进程同步", page=10),
        StructuredBlock(block_type=BlockType.PARAGRAPH, text="临界区是访问共享资源的代码段。", page=10),
        StructuredBlock(
            block_type=BlockType.TABLE,
            text="| P操作 | V操作 |\n| --- | --- |\n| wait | signal |",
            page=11,
        ),
        StructuredBlock(
            block_type=BlockType.CODE,
            text="sem_wait(&empty);\nsem_post(&full);",
            page=12,
            language="c",
        ),
        StructuredBlock(block_type=BlockType.FORMULA, text="S = S - 1", page=13),
    ]
    class _Doc:
        metadata_json = {"blocks": [b.to_dict() for b in blocks]}
        content = ""

    chunks = service._chunks_for_document(_Doc())
    texts = [t for t, _ in chunks]
    metas = {m.get("block_type"): (t, m) for t, m in chunks}
    assert "| P操作 | V操作 |" in "".join(texts)
    assert "sem_wait(&empty);" in "".join(texts)
    assert "S = S - 1" in "".join(texts)
    assert metas["table"][1]["page"] == 11
    assert metas["code"][1]["language"] == "c"


def test_rapid_table_html_to_markdown():
    """RapidTable HTML → Markdown 表格转换。"""
    html = "<table><tr><td>a</td><td>b</td></tr><tr><td>1</td><td>2</td></tr></table>"
    md = _html_to_markdown(html)
    assert "| a | b |" in md
    assert "| 1 | 2 |" in md
    assert "---" in md
